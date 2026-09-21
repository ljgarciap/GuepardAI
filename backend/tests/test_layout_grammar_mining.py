"""
test_layout_grammar_mining.py — Artistic Generation Engine v2, Phase 0.
docs/specs/artistic-generation-v2.md

Covers:
  - Additive geometry fields on TextTarget/TextSlot (left/x_pct/y_pct/w_pct/h_pct)
    — must not change any existing Template Merge behavior, only add fields.
  - services/generation/layout_grammar_service.py: PPTX region extraction (via
    the existing analyzer, real reuse), PDF region extraction (new, PyMuPDF,
    synthetic in-memory PDFs — no fixture files needed), and deterministic
    clustering.
  - agents/mine_layout_grammar.py::MineLayoutGrammarTool — mocked LLM, real
    db_session, persists/updates BrandLayoutGrammar.
"""
import io
from unittest.mock import patch

import pytest
from pptx.util import Inches

from services.templates.template_analyzer import analyze_template, _geometry_pct
from services.templates.template_config import TemplateMergeConfig
from services.templates.template_traversal import TextTarget, collect_text_targets
from services.generation.layout_grammar_service import (
    LayoutMiningConfig,
    MinedRegion,
    cluster_regions,
    extract_pdf_regions,
    extract_pptx_regions,
)
from agents.mine_layout_grammar import MineLayoutGrammarTool


def _blank_slide():
    from pptx import Presentation
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


# ---------------------------------------------------------------------------
# Additive geometry — TextTarget.left / TextSlot.x_pct etc.
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestAdditiveGeometry:

    def test_walk_populates_left_for_plain_textbox(self):
        _, slide = _blank_slide()
        slide.shapes.add_textbox(Inches(2), Inches(1), Inches(3), Inches(1))
        targets, _ = collect_text_targets(slide)
        assert targets[0].left == Inches(2)

    def test_geometry_pct_computes_percent_of_slide(self):
        target = TextTarget(
            key="1", text_frame=None, name="t", kind="shape", is_placeholder=False,
            left=Inches(2), top=Inches(1), width=Inches(3), height=Inches(1),
        )
        # Default python-pptx blank slide is 10in x 7.5in
        x_pct, y_pct, w_pct, h_pct = _geometry_pct(target, Inches(10), Inches(7.5))
        assert x_pct == pytest.approx(20.0, abs=0.1)
        assert y_pct == pytest.approx(13.33, abs=0.1)
        assert w_pct == pytest.approx(30.0, abs=0.1)
        assert h_pct == pytest.approx(13.33, abs=0.1)

    def test_geometry_pct_none_when_slide_dims_missing(self):
        target = TextTarget(key="1", text_frame=None, name="t", kind="shape", is_placeholder=False)
        assert _geometry_pct(target, 0, 0) == (None, None, None, None)

    def test_analyze_template_slots_carry_geometry(self, tmp_path):
        prs = _presentation_with_one_textbox("Hello world, this is a real slot")
        path = tmp_path / "t.pptx"
        prs.save(str(path))
        profiles = analyze_template(str(path), TemplateMergeConfig())
        assert profiles[0].slots[0].x_pct is not None
        assert 0 <= profiles[0].slots[0].x_pct <= 100

    def test_table_cell_left_is_cumulative_column_width(self):
        _, slide = _blank_slide()
        table_shape = slide.shapes.add_table(rows=1, cols=3, left=Inches(1), top=Inches(1),
                                              width=Inches(6), height=Inches(1))
        table = table_shape.table
        for col in table.columns:
            col.width = Inches(2)
        targets, _ = collect_text_targets(slide)
        by_key = {t.key: t for t in targets}
        # 3 columns of 2in each, table starts at left=Inches(1)
        assert by_key[f"{table_shape.shape_id}:r0c0"].left == Inches(1)
        assert by_key[f"{table_shape.shape_id}:r0c1"].left == Inches(3)
        assert by_key[f"{table_shape.shape_id}:r0c2"].left == Inches(5)


def _presentation_with_one_textbox(text: str):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = text
    return prs


# ---------------------------------------------------------------------------
# layout_grammar_service — extraction
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestExtractPptxRegions:

    def test_extracts_text_regions_with_percent_geometry(self, tmp_path):
        prs = _presentation_with_one_textbox("A real slot with enough text to survive filtering")
        path = tmp_path / "deck.pptx"
        prs.save(str(path))

        regions = extract_pptx_regions(str(path))
        assert len(regions) == 1
        assert regions[0].region_type == "text"
        assert regions[0].page_idx == 0
        assert 0 <= regions[0].x_pct <= 100


@pytest.mark.unit
class TestExtractPdfRegions:

    def _pdf_bytes_with(self, text_at=(36, 40), rect=None, image_rect=None):
        import fitz
        doc = fitz.open()
        page = doc.new_page(width=720, height=540)
        if text_at:
            page.insert_text(text_at, "Real Title Text Here", fontsize=24)
        if rect:
            page.draw_rect(fitz.Rect(*rect), color=(1, 0, 0), fill=(1, 0, 0))
        if image_rect:
            from PIL import Image
            buf = io.BytesIO()
            Image.new("RGB", (100, 60), color=(0, 128, 255)).save(buf, format="PNG")
            buf.seek(0)
            page.insert_image(fitz.Rect(*image_rect), stream=buf.read())
        data = doc.tobytes()
        doc.close()
        return data

    def test_extracts_text_region(self, tmp_path):
        pdf_path = tmp_path / "a.pdf"
        pdf_path.write_bytes(self._pdf_bytes_with())
        regions = extract_pdf_regions(str(pdf_path))
        text_regions = [r for r in regions if r.region_type == "text"]
        assert len(text_regions) == 1
        assert "Real Title" in text_regions[0].text_hint

    def test_extracts_shape_region_with_color(self, tmp_path):
        pdf_path = tmp_path / "b.pdf"
        pdf_path.write_bytes(self._pdf_bytes_with(text_at=None, rect=(50, 100, 150, 200)))
        regions = extract_pdf_regions(str(pdf_path))
        shape_regions = [r for r in regions if r.region_type == "shape"]
        assert len(shape_regions) == 1
        assert shape_regions[0].color == "#FF0000"

    def test_rotated_page_geometry_stays_in_bounds_and_reorients(self, tmp_path):
        # Real-world regression: a 90°-rotated page (confirmed on the actual
        # Embonor Insumos sample) has get_text/get_drawings return coordinates
        # in the page's RAW/mediabox space, not its display (page.rect) space —
        # normalizing by page.rect without correcting for rotation produced
        # percentages up to ~130%, and silently transposed top<->left.
        import fitz
        doc = fitz.open()
        # mediabox 400x600 portrait, rotated 90 -> displays as 600x400 landscape
        page = doc.new_page(width=400, height=600)
        page.set_rotation(90)
        # A shape in the raw mediabox's bottom-left quadrant: (0,300)-(200,600),
        # i.e. raw x in [0,50%] and raw y in [50%,100%] of the 400x600 mediabox.
        # Kept well under the 85% max-area noise filter (unlike a full-bleed
        # background) so this isolates the rotation math, not the area filter.
        page.draw_rect(fitz.Rect(0, 300, 200, 600), color=(0, 1, 0), fill=(0, 1, 0))
        pdf_path = tmp_path / "rotated.pdf"
        pdf_path.write_bytes(doc.tobytes())
        doc.close()

        regions = extract_pdf_regions(str(pdf_path))
        shapes = [r for r in regions if r.region_type == "shape"]
        assert len(shapes) == 1
        shape = shapes[0]
        assert 0 <= shape.x_pct <= 100
        assert 0 <= shape.y_pct <= 100
        assert 0 <= shape.w_pct <= 100
        assert 0 <= shape.h_pct <= 100

    def test_extracts_image_region(self, tmp_path):
        pdf_path = tmp_path / "c.pdf"
        pdf_path.write_bytes(self._pdf_bytes_with(text_at=None, image_rect=(200, 200, 400, 320)))
        regions = extract_pdf_regions(str(pdf_path))
        image_regions = [r for r in regions if r.region_type == "image"]
        assert len(image_regions) == 1

    def test_tiny_shape_filtered_as_noise(self, tmp_path):
        pdf_path = tmp_path / "d.pdf"
        # 2x2pt speck on a 720x540 page is far below the default 0.3% area threshold
        pdf_path.write_bytes(self._pdf_bytes_with(text_at=None, rect=(0, 0, 2, 2)))
        regions = extract_pdf_regions(str(pdf_path), LayoutMiningConfig())
        assert not [r for r in regions if r.region_type == "shape"]

    def test_near_white_fill_filtered(self, tmp_path):
        import fitz
        doc = fitz.open()
        page = doc.new_page(width=720, height=540)
        page.draw_rect(fitz.Rect(50, 100, 150, 200), color=(1, 1, 1), fill=(0.99, 0.99, 0.99))
        pdf_path = tmp_path / "e.pdf"
        pdf_path.write_bytes(doc.tobytes())
        doc.close()
        regions = extract_pdf_regions(str(pdf_path))
        assert not [r for r in regions if r.region_type == "shape"]


# ---------------------------------------------------------------------------
# layout_grammar_service — clustering
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestClusterRegions:

    def _region(self, page_idx, x, y, w=16.0, h=16.0, region_type="shape", role=None):
        return MinedRegion(page_idx=page_idx, region_type=region_type, x_pct=x, y_pct=y, w_pct=w, h_pct=h, role=role)

    def test_pages_with_matching_geometry_join_one_cluster(self):
        # All 3 land in the same rounding bucket (tolerance=4.0 default,
        # bucket centered on 40.0/28.0) — kept well clear of the bucket edge
        # (42.0) so the test isn't sensitive to round-half-to-even behavior.
        regions = [
            self._region(4, 41.0, 28.0),
            self._region(9, 40.5, 28.5),
            self._region(14, 41.5, 27.5),
        ]
        clusters = cluster_regions(regions)
        assert len(clusters) == 1
        assert clusters[0]["page_indices"] == [4, 9, 14]

    def test_pages_with_different_geometry_form_separate_clusters(self):
        regions = [
            self._region(0, 42.0, 28.0),
            self._region(1, 5.0, 90.0),   # far outside tolerance
        ]
        clusters = cluster_regions(regions)
        assert len(clusters) == 2

    def test_single_page_cluster_still_produced(self):
        regions = [self._region(0, 0.0, 0.0, w=100.0, h=22.0, region_type="image", role=None)]
        clusters = cluster_regions(regions)
        assert len(clusters) == 1
        assert clusters[0]["page_indices"] == [0]

    def test_shared_slots_carries_role_and_geometry(self):
        regions = [self._region(0, 42.0, 28.0, region_type="text", role="title")]
        clusters = cluster_regions(regions)
        slot = clusters[0]["shared_slots"][0]
        assert slot["role"] == "title"
        assert slot["x_pct"] == 42.0


# ---------------------------------------------------------------------------
# MineLayoutGrammarTool
# ---------------------------------------------------------------------------

@pytest.mark.unit
class TestMineLayoutGrammarTool:

    def test_empty_clusters_returns_empty_without_llm_call(self, mock_llm_calls):
        tool = MineLayoutGrammarTool()
        result = tool.run(brand_id=1, source_filename="x.pptx", clusters=[])
        assert result == {"signatures": []}
        mock_llm_calls["generate_json"].assert_not_called()

    def test_persists_new_brand_layout_grammar_row(self, mock_llm_calls, db_session, sample_brand):
        import models
        mock_llm_calls["generate_json"].return_value = {
            "signatures": [{
                "name": "donut-kpi-trio", "source_slide_indices": [4, 9, 14],
                "content_shape": "metric_comparison", "slots": [], "motifs": ["ribbon_band"],
                "confidence": 0.85,
            }]
        }
        tool = MineLayoutGrammarTool()
        clusters = [{"page_indices": [4, 9, 14], "shared_slots": []}]
        # Capture the id before run(): MineLayoutGrammarTool's `finally: db.close()`
        # detaches sample_brand once SessionLocal is patched to the fixture session.
        brand_id = sample_brand.id
        # MineLayoutGrammarTool opens its own SessionLocal() — patch it to the
        # fixture's session (same pattern as test_render_agent_canvas_elements_
        # wiring.py) so it sees sample_brand's flushed-not-committed row.
        with patch("agents.mine_layout_grammar.SessionLocal", return_value=db_session):
            result = tool.run(brand_id=brand_id, source_filename="embonor.pdf", clusters=clusters)

        assert result["signatures"][0]["name"] == "donut-kpi-trio"
        row = db_session.query(models.BrandLayoutGrammar).filter(
            models.BrandLayoutGrammar.brand_id == brand_id,
            models.BrandLayoutGrammar.source_filename == "embonor.pdf",
        ).first()
        assert row is not None
        assert row.signatures_json[0]["name"] == "donut-kpi-trio"
        assert row.raw_extraction == clusters

    def test_rerun_updates_existing_row_not_duplicate(self, mock_llm_calls, db_session, sample_brand):
        import models
        tool = MineLayoutGrammarTool()
        brand_id = sample_brand.id

        with patch("agents.mine_layout_grammar.SessionLocal", return_value=db_session):
            mock_llm_calls["generate_json"].return_value = {"signatures": [{"name": "first-pass"}]}
            tool.run(brand_id=brand_id, source_filename="deck.pptx", clusters=[{"page_indices": [0], "shared_slots": []}])

            mock_llm_calls["generate_json"].return_value = {"signatures": [{"name": "second-pass"}]}
            tool.run(brand_id=brand_id, source_filename="deck.pptx", clusters=[{"page_indices": [0, 1], "shared_slots": []}])

        rows = db_session.query(models.BrandLayoutGrammar).filter(
            models.BrandLayoutGrammar.brand_id == brand_id,
            models.BrandLayoutGrammar.source_filename == "deck.pptx",
        ).all()
        assert len(rows) == 1
        assert rows[0].signatures_json[0]["name"] == "second-pass"
